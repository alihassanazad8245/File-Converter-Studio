"""Tests for each format converter and the conversion engine."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from markdown_converter.config import ConversionOptions
from markdown_converter.engine import ConversionRequest, convert_batch, convert_one, resolve_and_convert
from markdown_converter.parser import parse_markdown

SAMPLE_MD = """# Sample Report

This is **bold**, *italic*, and `inline code`.

## Section

- item one
- item two

| Name | Score |
|------|-------|
| Ali  | 95    |

```python
print("hi")
```

> A quote.

[a link](https://example.com)
"""


def _write_sample(tmp_path: Path, name: str = "sample.md") -> Path:
    path = tmp_path / name
    path.write_text(SAMPLE_MD, encoding="utf-8")
    return path


# --------------------------------------------------------------------------- #
# Individual converters via convert_one
# --------------------------------------------------------------------------- #


def test_html_converter_produces_valid_standalone_page(tmp_path: Path) -> None:
    src = _write_sample(tmp_path)
    out = tmp_path / "out.html"
    result = convert_one(src, "html", out, ConversionOptions())
    assert result.success
    content = out.read_text(encoding="utf-8")
    assert "<html" in content and "Sample Report" in content


def test_pdf_converter_produces_nonempty_file(tmp_path: Path) -> None:
    src = _write_sample(tmp_path)
    out = tmp_path / "out.pdf"
    result = convert_one(src, "pdf", out, ConversionOptions())
    assert result.success
    assert out.stat().st_size > 500
    assert out.read_bytes().startswith(b"%PDF")


def test_docx_converter_produces_valid_document(tmp_path: Path) -> None:
    src = _write_sample(tmp_path)
    out = tmp_path / "out.docx"
    result = convert_one(src, "docx", out, ConversionOptions())
    assert result.success
    doc = Document(str(out))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Sample Report" in all_text
    assert len(doc.tables) == 1


def test_txt_converter_strips_formatting(tmp_path: Path) -> None:
    src = _write_sample(tmp_path)
    out = tmp_path / "out.txt"
    result = convert_one(src, "txt", out, ConversionOptions())
    assert result.success
    content = out.read_text(encoding="utf-8")
    assert "SAMPLE REPORT" in content  # headings upper-cased
    assert "**" not in content  # markdown syntax stripped


def test_md_converter_normalises_line_endings(tmp_path: Path) -> None:
    src = tmp_path / "crlf.md"
    src.write_bytes(b"# Title\r\n\r\nBody\r\n")
    out = tmp_path / "out.md"
    result = convert_one(src, "md", out, ConversionOptions())
    assert result.success
    assert "\r" not in out.read_text(encoding="utf-8")


def test_json_converter_produces_structured_tree(tmp_path: Path) -> None:
    src = _write_sample(tmp_path)
    out = tmp_path / "out.json"
    result = convert_one(src, "json", out, ConversionOptions())
    assert result.success
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["title"] == "Sample Report"
    types = [n["type"] for n in data["nodes"]]
    assert "heading" in types
    assert "table" in types
    assert "code_block" in types
    assert any(link["href"] == "https://example.com" for link in data["links"])


def test_convert_one_rejects_unsupported_format(tmp_path: Path) -> None:
    from markdown_converter.errors import ConversionError

    src = _write_sample(tmp_path)
    with pytest.raises(ConversionError):
        convert_one(src, "epub", tmp_path / "out.epub", ConversionOptions())


# --------------------------------------------------------------------------- #
# Missing image handling across formats (must never crash)
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize("fmt", ["pdf", "docx", "html"])
def test_missing_image_produces_warning_not_crash(tmp_path: Path, fmt: str) -> None:
    src = tmp_path / "with_image.md"
    src.write_text("# Doc\n\n![missing](nope.png)\n", encoding="utf-8")
    out = tmp_path / f"out.{fmt}"
    result = convert_one(src, fmt, out, ConversionOptions())
    assert result.success  # the rest of the document still converts
    assert any("nope.png" in w for w in result.warnings)


# --------------------------------------------------------------------------- #
# Engine: full resolve + convert
# --------------------------------------------------------------------------- #


def test_resolve_and_convert_same_folder(tmp_path: Path) -> None:
    src = _write_sample(tmp_path)
    request = ConversionRequest(source_path=src, formats=["html"], same_folder=True)
    results = resolve_and_convert(request)
    assert len(results) == 1
    assert Path(results[0].output_path).parent == tmp_path


def test_resolve_and_convert_custom_folder(tmp_path: Path) -> None:
    src = _write_sample(tmp_path)
    other = tmp_path / "converted"
    request = ConversionRequest(source_path=src, formats=["html"], same_folder=False,
                                custom_output_dir=str(other))
    results = resolve_and_convert(request)
    assert Path(results[0].output_path).parent == other
    assert other.exists()


def test_resolve_and_convert_multiple_formats(tmp_path: Path) -> None:
    src = _write_sample(tmp_path)
    request = ConversionRequest(source_path=src, formats=["html", "txt", "json"], same_folder=True)
    results = resolve_and_convert(request)
    assert len(results) == 3
    assert all(r.success for r in results)


def test_resolve_and_convert_custom_filename(tmp_path: Path) -> None:
    src = _write_sample(tmp_path)
    request = ConversionRequest(source_path=src, formats=["pdf"], same_folder=True,
                                custom_filename="My Report")
    results = resolve_and_convert(request)
    assert Path(results[0].output_path).name == "My Report.pdf"


def test_resolve_and_convert_invalid_source_raises() -> None:
    from markdown_converter.errors import FileValidationError

    request = ConversionRequest(source_path=Path("/nonexistent/file.md"), formats=["pdf"])
    with pytest.raises(FileValidationError):
        resolve_and_convert(request)


def test_resolve_and_convert_replace_collision(tmp_path: Path) -> None:
    src = _write_sample(tmp_path)
    existing = tmp_path / "sample.html"
    existing.write_text("old content")
    request = ConversionRequest(source_path=src, formats=["html"], same_folder=True)
    results = resolve_and_convert(request, on_collision="replace")
    assert "old content" not in existing.read_text(encoding="utf-8")


def test_resolve_and_convert_new_copy_collision(tmp_path: Path) -> None:
    src = _write_sample(tmp_path)
    existing = tmp_path / "sample.html"
    existing.write_text("old content")
    request = ConversionRequest(source_path=src, formats=["html"], same_folder=True)
    results = resolve_and_convert(request, on_collision="new_copy")
    assert Path(results[0].output_path).name == "sample (1).html"
    assert existing.read_text(encoding="utf-8") == "old content"  # untouched


def test_resolve_and_convert_cancel_collision_skips(tmp_path: Path) -> None:
    src = _write_sample(tmp_path)
    existing = tmp_path / "sample.html"
    existing.write_text("old content")
    request = ConversionRequest(source_path=src, formats=["html"], same_folder=True)
    results = resolve_and_convert(request, on_collision="cancel")
    assert results[0].success is False
    assert "already exists" in results[0].message


# --------------------------------------------------------------------------- #
# Batch conversion
# --------------------------------------------------------------------------- #


def test_convert_batch_processes_multiple_files(tmp_path: Path) -> None:
    src1 = _write_sample(tmp_path, "one.md")
    src2 = _write_sample(tmp_path, "two.md")
    requests = [
        ConversionRequest(source_path=src1, formats=["html"], same_folder=True),
        ConversionRequest(source_path=src2, formats=["html"], same_folder=True),
    ]
    results = convert_batch(requests)
    assert len(results) == 2
    assert all(r.success for r in results)


def test_convert_batch_one_bad_file_does_not_abort_others(tmp_path: Path) -> None:
    good = _write_sample(tmp_path, "good.md")
    requests = [
        ConversionRequest(source_path=Path("/nonexistent.md"), formats=["html"]),
        ConversionRequest(source_path=good, formats=["html"], same_folder=True),
    ]
    results = convert_batch(requests)
    assert len(results) == 2
    assert results[0].success is False
    assert results[1].success is True


def test_convert_batch_reports_progress(tmp_path: Path) -> None:
    src1 = _write_sample(tmp_path, "one.md")
    src2 = _write_sample(tmp_path, "two.md")
    requests = [
        ConversionRequest(source_path=src1, formats=["html"], same_folder=True),
        ConversionRequest(source_path=src2, formats=["html"], same_folder=True),
    ]
    seen = []
    convert_batch(requests, progress_callback=lambda i, t, p: seen.append((i, t)))
    assert seen == [(0, 2), (1, 2)]
