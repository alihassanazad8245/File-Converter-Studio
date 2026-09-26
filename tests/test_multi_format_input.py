"""Tests for load_any_document - the multi-format input loader."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from markdown_converter.parser import load_any_document


def test_load_any_document_markdown(tmp_path: Path) -> None:
    path = tmp_path / "doc.md"
    path.write_text("# Title\n\nHello **world**.")
    doc = load_any_document(path)
    assert doc.title == "Title"
    assert doc.soup.find("strong") is not None


def test_load_any_document_txt_splits_paragraphs(tmp_path: Path) -> None:
    path = tmp_path / "doc.txt"
    path.write_text("First paragraph.\n\nSecond paragraph.")
    doc = load_any_document(path)
    paragraphs = doc.soup.find_all("p")
    assert len(paragraphs) == 2
    assert "First paragraph" in paragraphs[0].get_text()


def test_load_any_document_html_extracts_body(tmp_path: Path) -> None:
    path = tmp_path / "doc.html"
    path.write_text("<html><head><title>My Page</title></head><body><h1>Hi</h1></body></html>")
    doc = load_any_document(path)
    assert doc.title == "My Page"
    assert doc.soup.find("h1").get_text() == "Hi"


def test_load_any_document_docx(tmp_path: Path) -> None:
    path = tmp_path / "doc.docx"
    document = Document()
    document.add_heading("Report Title", level=1)
    document.add_paragraph("Some content.")
    document.save(str(path))

    doc = load_any_document(path)
    assert doc.soup.find("h1") is not None
    assert "Some content" in doc.soup.get_text()


def test_load_any_document_rejects_unsupported_extension(tmp_path: Path) -> None:
    path = tmp_path / "doc.xyz"
    path.write_text("content")
    with pytest.raises(ValueError):
        load_any_document(path)
