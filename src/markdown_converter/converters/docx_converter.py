"""Markdown -> DOCX converter.

Walks the same rendered HTML tree used by every other format and builds a
real python-docx document element by element (headings, paragraphs with
inline bold/italic/code runs, lists, tables, code blocks, blockquotes,
images, links). This is a genuine structural conversion, not a wrapped
plain-text dump.
"""

from __future__ import annotations

import time
from pathlib import Path

from bs4.element import NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from ..config import ConversionOptions
from ..models import ConversionResult
from ..parser import ParsedDocument, resolve_local_path
from .base import BaseConverter

_CODE_BG = RGBColor(0xF2, 0xF5, 0xF8)
_LINK_COLOR = RGBColor(0x25, 0x63, 0xEB)


class DocxConverter(BaseConverter):
    format_name = "docx"

    def convert(self, doc: ParsedDocument, output_path: Path, options: ConversionOptions) -> ConversionResult:
        start = time.monotonic()
        opts = options.docx
        document = Document()

        style = document.styles["Normal"]
        style.font.name = opts.font_family
        style.font.size = Pt(opts.font_size)

        missing_images: list[str] = []

        if opts.include_toc:
            headings = doc.soup.find_all(["h1", "h2", "h3"])
            if headings:
                document.add_heading("Table of Contents", level=1)
                for h in headings:
                    document.add_paragraph(h.get_text(strip=True), style="List Bullet")
                document.add_page_break()

        for element in doc.soup.find_all(recursive=False):
            self._render_block(document, element, doc.source_dir, missing_images)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_path))

        warnings = [f"Image not found: {src}" for src in missing_images]
        return ConversionResult(
            source_path=str(doc.source_dir), output_path=str(output_path), format="docx",
            success=True, message="DOCX created successfully.", warnings=warnings,
            duration_seconds=round(time.monotonic() - start, 3),
        )

    # ------------------------------------------------------------------ #
    # Block-level rendering
    # ------------------------------------------------------------------ #

    def _render_block(self, document: Document, element, source_dir: Path, missing: list[str]) -> None:
        if isinstance(element, NavigableString):
            text = str(element).strip()
            if text:
                document.add_paragraph(text)
            return
        if not isinstance(element, Tag):
            return

        name = element.name
        if name and len(name) == 2 and name[0] == "h" and name[1].isdigit():
            level = min(int(name[1]), 9)
            document.add_heading(element.get_text(strip=True), level=level)
        elif name == "p":
            paragraph = document.add_paragraph()
            self._render_inline(paragraph, element, source_dir, missing)
        elif name in ("ul", "ol"):
            self._render_list(document, element, ordered=(name == "ol"))
        elif name == "table":
            self._render_table(document, element)
        elif name in ("pre",):
            self._render_code_block(document, element)
        elif name == "blockquote":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(24)
            run = paragraph.add_run(element.get_text(strip=True))
            run.italic = True
        elif name == "hr":
            document.add_paragraph("―" * 40)
        elif name == "img":
            self._render_image(document, element, source_dir, missing)
        elif name in ("div", "section", "article"):
            for child in element.find_all(recursive=False):
                self._render_block(document, child, source_dir, missing)
        else:
            text = element.get_text(strip=True)
            if text:
                document.add_paragraph(text)

    def _render_list(self, document: Document, element: Tag, ordered: bool, depth: int = 0) -> None:
        style = "List Number" if ordered else "List Bullet"
        for item in element.find_all("li", recursive=False):
            nested = item.find(["ul", "ol"], recursive=False)
            text_parts = [c for c in item.children if not (isinstance(c, Tag) and c.name in ("ul", "ol"))]
            paragraph = document.add_paragraph(style=style)
            for part in text_parts:
                if isinstance(part, NavigableString):
                    paragraph.add_run(str(part))
                elif isinstance(part, Tag):
                    self._render_inline_into(paragraph, part)
            if nested:
                self._render_list(document, nested, ordered=(nested.name == "ol"), depth=depth + 1)

    def _render_table(self, document: Document, element: Tag) -> None:
        rows = element.find_all("tr")
        if not rows:
            return
        n_cols = max(len(r.find_all(["td", "th"])) for r in rows)
        table = document.add_table(rows=0, cols=n_cols)
        table.style = "Light Grid Accent 1"
        for row in rows:
            cells = row.find_all(["td", "th"])
            docx_row = table.add_row().cells
            for i, cell in enumerate(cells):
                if i >= n_cols:
                    break
                docx_row[i].text = cell.get_text(strip=True)
                if cell.name == "th":
                    for p in docx_row[i].paragraphs:
                        for r in p.runs:
                            r.bold = True

    def _render_code_block(self, document: Document, element: Tag) -> None:
        text = element.get_text()
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(12)
        run = paragraph.add_run(text.rstrip("\n"))
        run.font.name = "Courier New"
        run.font.size = Pt(9.5)

    def _render_image(self, document: Document, element: Tag, source_dir: Path, missing: list[str]) -> None:
        src = element.get("src", "")
        resolved = resolve_local_path(src, source_dir)
        if resolved is None or not resolved.is_file():
            missing.append(src)
            document.add_paragraph(f"[Image not found: {src}]")
            return
        try:
            document.add_picture(str(resolved), width=Pt(360))
        except Exception:  # noqa: BLE001 - unsupported image format, degrade gracefully
            missing.append(src)
            document.add_paragraph(f"[Could not embed image: {src}]")

    # ------------------------------------------------------------------ #
    # Inline rendering (bold, italic, code, links) within a paragraph
    # ------------------------------------------------------------------ #

    def _render_inline(self, paragraph, element: Tag, source_dir: Path, missing: list[str]) -> None:
        for child in element.children:
            if isinstance(child, NavigableString):
                paragraph.add_run(str(child))
            elif isinstance(child, Tag):
                if child.name == "img":
                    self._render_image_inline(paragraph, child, source_dir, missing)
                else:
                    self._render_inline_into(paragraph, child)

    def _render_inline_into(self, paragraph, element: Tag) -> None:
        """Render one inline element (and its children) as runs on ``paragraph``."""
        bold = element.name in ("strong", "b")
        italic = element.name in ("em", "i")
        code = element.name == "code"
        strike = element.name in ("del", "s")
        is_link = element.name == "a"

        if element.name in ("strong", "b", "em", "i", "code", "del", "s", "a"):
            for grandchild in element.children:
                if isinstance(grandchild, NavigableString):
                    run = paragraph.add_run(str(grandchild))
                    run.bold = bold
                    run.italic = italic
                    run.font.strike = strike
                    if code:
                        run.font.name = "Courier New"
                        run.font.highlight_color = None
                    if is_link:
                        run.font.color.rgb = _LINK_COLOR
                        run.underline = True
                elif isinstance(grandchild, Tag):
                    self._render_inline_into(paragraph, grandchild)
        else:
            text = element.get_text()
            if text:
                paragraph.add_run(text)

    def _render_image_inline(self, paragraph, element: Tag, source_dir: Path, missing: list[str]) -> None:
        src = element.get("src", "")
        resolved = resolve_local_path(src, source_dir)
        if resolved is None or not resolved.is_file():
            missing.append(src)
            paragraph.add_run(f"[Image not found: {src}]")
            return
        try:
            run = paragraph.add_run()
            run.add_picture(str(resolved), width=Pt(300))
        except Exception:  # noqa: BLE001
            missing.append(src)
            paragraph.add_run(f"[Could not embed image: {src}]")
